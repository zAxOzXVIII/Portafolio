"""Genera PDF del CV desde el documento Word fuente."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
DOCX_SOURCE = Path(r"d:\Documentos\Word\Marden B. curriculum.docx")
PDF_OUTPUT = ROOT / "assets" / "cv" / "Marden_Barrera_CV.pdf"
FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


class CVPdf(FPDF):
    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Arial", "", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, f"Página {self.page_no()}", align="C")


def style_for(paragraph_style: str) -> tuple[str, int, bool]:
    if paragraph_style == "Title":
        return ("Arial", 18, True)
    if paragraph_style == "Heading 1":
        return ("Arial", 12, True)
    if paragraph_style == "Heading 2":
        return ("Arial", 10, True)
    if paragraph_style == "Subtitle":
        return ("Arial", 9, False)
    return ("Arial", 10, False)


def build_pdf() -> None:
    doc = Document(DOCX_SOURCE)
    pdf = CVPdf()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()

    pdf.add_font("Arial", "", str(FONT_REGULAR))
    pdf.add_font("Arial", "B", str(FONT_BOLD))

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            pdf.ln(2)
            continue

        family, size, bold = style_for(para.style.name)
        pdf.set_font(family, "B" if bold else "", size)

        if para.style.name == "Title":
            pdf.set_text_color(20, 20, 20)
            pdf.cell(0, 10, text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            continue

        if para.style.name == "Heading 1":
            pdf.set_text_color(30, 30, 30)
            pdf.ln(4)
            pdf.cell(0, 7, text.upper(), new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(200, 200, 200)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(3)
            continue

        if para.style.name in {"Heading 2", "Subtitle"}:
            pdf.set_text_color(45, 45, 45)
        else:
            pdf.set_text_color(55, 55, 55)

        bullet = para.style.name == "List Paragraph"
        line = f"• {text}" if bullet else text
        pdf.multi_cell(0, 5.5, line)
        pdf.ln(1)

    PDF_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(PDF_OUTPUT))
    print(f"PDF generado: {PDF_OUTPUT}")


if __name__ == "__main__":
    build_pdf()
